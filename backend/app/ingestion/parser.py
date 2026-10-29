import io
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Set, Tuple

from app.config import MAX_CHARS, MAX_PAGES
from app.models.schemas import PageSpan

_NUMBERED_LINE = re.compile(
    r"^\s*(?:\d{1,2}(?:\.\d{1,2}){0,3}\.?\s+\S|\([a-zA-Z0-9]{1,4}\)\s+\S|[A-Z]\.\s+\S|[ivxIVX]{1,5}\.\s+\S"
    r"|(?:ARTICLE|Article|SECTION|Section|SCHEDULE|Schedule|ANNEXURE|Annexure|APPENDIX|Appendix|EXHIBIT|Exhibit)\b"
    r"|[•\-–\*]\s+\S)"
)


# A table cell holding only a clause number, e.g. "1.", "2.3", "(a)", "iv." in two-column numbered layouts.
_MARKER_CELL = re.compile(r"^(?:\d{1,2}(?:\.\d{1,2}){0,3}\.?|\([a-zA-Z0-9]{1,4}\)|[a-zA-Z]\.|[ivxIVX]{1,5}\.)$")


class DocumentParseError(ValueError):
    pass


@dataclass
class IngestedDocument:
    text: str
    pages: List[PageSpan]
    source_type: str
    page_count: int
    ocr_used: bool = False
    heading_hints: List[int] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def normalize_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    text = text.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'").replace("\u2018", "'")
    text = text.replace("\u2014", " - ").replace("\u2013", "-").replace("\u2022", "•")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _is_heading_line(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 90:
        return False
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and sum(c.isupper() for c in letters) / len(letters) > 0.8


_TRAILING_FUNCTION_WORDS = {"a", "an", "and", "or", "the", "of", "to", "in", "if", "for", "with", "by", "that", "on", "at"}
_HEAVY_FONT = re.compile(r"bold|black|heavy|semibold|demi|medium", re.I)
_PAGE_NUMBER = re.compile(r"^(?:page\s+)?[-(]?\s*\d{1,4}\s*[-)]?(?:\s*(?:of|/)\s*\d{1,4})?$|^\d{1,4}\s*\|\s*page$", re.I)
_EDGE_LINES = 4


def _edge_key(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def strip_running_lines(page_texts: List[str]) -> List[str]:
    """Drop page numbers and letterhead/footer lines repeated at the top or bottom of most pages."""
    pages = [t.split("\n") for t in page_texts]
    edges: List[set] = []
    counts: dict = {}
    for lines in pages:
        idx = [i for i, ln in enumerate(lines) if ln.strip()]
        edge = set(idx[:_EDGE_LINES] + idx[-_EDGE_LINES:])
        edges.append(edge)
        for key in {_edge_key(lines[i]) for i in edge}:
            counts[key] = counts.get(key, 0) + 1
    threshold = max(2, -(-len(pages) // 2))
    cleaned = []
    for lines, edge in zip(pages, edges):
        nonblank = [i for i, ln in enumerate(lines) if ln.strip()]
        outer = set(nonblank[:2] + nonblank[-2:])
        drop = set()
        for i in edge:
            s = lines[i].strip()
            repeated = (len(pages) >= 2 and len(s) <= 80 and counts.get(_edge_key(s), 0) >= threshold
                        and sum(c.isalpha() for c in s) >= 3)
            if repeated or (i in outer and _PAGE_NUMBER.match(s)):
                drop.add(i)
        cleaned.append("\n".join(ln for i, ln in enumerate(lines) if i not in drop))
    return cleaned


def styled_lines(pages) -> Tuple[Set[str], Set[str]]:
    """Visual structure that plain text loses: heading lines (larger or bolder than the body text) and
    lines that open with a bold run-in label such as "Start Date: Your first day...".
    """
    lines: List[Tuple[str, float, bool]] = []
    run_ins: Set[str] = set()
    sizes: Dict[float, int] = {}
    bold_chars = total = 0
    for page in pages:
        for block in page.get_text("dict").get("blocks", []):
            block_lines = [ln for ln in block.get("lines", []) if any(s.get("text", "").strip() for s in ln.get("spans", []))]
            lead_paragraph = len(block_lines) >= 3
            for line in block_lines:
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                n = sum(len(s["text"].strip()) for s in spans)
                heavy = [bool(s["flags"] & 16 or _HEAVY_FONT.search(s["font"])) for s in spans]
                bold = all(heavy)
                text = "".join(s["text"] for s in spans)
                label = spans[0]["text"].strip()
                after = spans[1]["text"].lstrip()[:1] if len(spans) > 1 else ""
                if (heavy[0] and not bold and len(label) <= 40 and label[:1].isupper()
                        and (label[-1] in ":." or after in (":", "."))):
                    run_ins.add(normalize_text(text))
                if not lead_paragraph:
                    lines.append((text, max(s["size"] for s in spans), bold))
                for s in spans:
                    key = round(s["size"], 1)
                    sizes[key] = sizes.get(key, 0) + len(s["text"].strip())
                total += n
                bold_chars += n if bold else 0
    if not sizes:
        return set(), set()
    body = max(sizes, key=sizes.get)
    use_bold = bold_chars < total * 0.4
    found = set()
    for raw, size, bold in lines:
        text = normalize_text(raw)
        if not 3 <= len(text) <= 100 or text[-1] in ",;" or text[0] in "•-–*" or sum(c.isalpha() for c in text) < 3:
            continue
        words = text.split()
        if text[0].islower() or words[-1].lower() in _TRAILING_FUNCTION_WORDS or (text.endswith(".") and len(words) > 6):
            continue
        if size >= body + 1.0 or (use_bold and bold and size >= body - 0.5):
            found.add(text)
    return found, (run_ins if use_bold else set())


def reflow_pdf_lines(page_text: str, headings: Set[str] = frozenset(), run_ins: Set[str] = frozenset()) -> str:
    """Join hard-wrapped PDF lines into paragraphs while keeping headings and numbered items on their own lines."""
    lines = [ln.rstrip() for ln in page_text.split("\n")]
    widths = sorted(len(ln) for ln in lines if ln.strip())
    typical = widths[int(len(widths) * 0.75)] if widths else 80
    caps = [_is_heading_line(ln) for ln in lines]
    paragraphs: List[str] = []
    current = ""
    marker = ""
    for i, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            if current and not marker:
                paragraphs.append(current)
                current = ""
            continue
        if _MARKER_CELL.match(stripped):
            if current:
                paragraphs.append(current)
                current = ""
            marker = f"{marker} {stripped}".strip()
            continue
        key = normalize_text(stripped) if headings or run_ins else ""
        styled, run_in = key in headings, key in run_ins
        if marker:
            stripped, marker = f"{marker} {stripped}", ""
        # A wrapped ALL-CAPS paragraph (arbitration notices) is body text, not a run of headings.
        caps_body = caps[i] and len(stripped) >= typical * 0.6 and (
            (i + 1 < len(lines) and caps[i + 1]) or (i > 0 and caps[i - 1] and bool(current)))
        heading = (_is_heading_line(stripped) or styled) and not caps_body
        caps_start = caps_body and not (i > 0 and caps[i - 1])
        starts_new = bool(_NUMBERED_LINE.match(stripped)) or heading or caps_start or run_in
        if current and starts_new:
            paragraphs.append(current)
            current = ""
        if not current:
            current = stripped
        elif current.endswith("-") and not current.endswith(" -"):
            current = current[:-1] + stripped
        else:
            current = f"{current} {stripped}"
        short_line = len(stripped) < typical * 0.6
        if heading or (short_line and stripped[-1:] in ".:;"):
            paragraphs.append(current)
            current = ""
    if marker:
        current = f"{current} {marker}".strip() if current else marker
    if current:
        paragraphs.append(current)
    return "\n\n".join(paragraphs)


def _ocr_page(page) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        pix = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(image)
    except Exception:
        return ""


def _continues_across_pages(previous: str, following: str) -> bool:
    """A sentence cut by a page break: no closing punctuation, and the next page carries on mid-sentence."""
    tail, head = previous.rstrip(), following.lstrip()
    if not tail or not head or tail[-1] in '.:;!?)"\'' or _NUMBERED_LINE.match(head):
        return False
    first_line = head.split("\n", 1)[0]
    last_line = tail.rsplit("\n", 1)[-1]
    return head[0].islower() or (_is_heading_line(last_line[-90:]) and len(last_line) > 90 and _is_heading_line(first_line[:90]))


def _join_pages(page_texts: List[str]) -> Tuple[str, List[PageSpan]]:
    parts: List[str] = []
    pages: List[PageSpan] = []
    cursor = 0
    for idx, raw in enumerate(page_texts, start=1):
        page_text = normalize_text(raw)
        if parts:
            sep = " " if _continues_across_pages(parts[-1], page_text) else "\n\n"
            parts.append(sep)
            cursor += len(sep)
        start = cursor
        parts.append(page_text)
        cursor += len(page_text)
        pages.append(PageSpan(page=idx, start=start, end=cursor))
    return "".join(parts), pages


def parse_pdf(data: bytes) -> IngestedDocument:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise DocumentParseError("PDF support requires PyMuPDF (pip install pymupdf).") from exc
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentParseError(f"Could not open PDF: {exc}") from exc
    if doc.needs_pass:
        raise DocumentParseError("This PDF is password-protected. Remove the password and upload again.")

    warnings: List[str] = []
    page_count = doc.page_count
    if page_count > MAX_PAGES:
        warnings.append(f"Only the first {MAX_PAGES} of {page_count} pages were analyzed.")
    raw_pages: List[str] = []
    ocr_used = False
    empty_pages = 0
    for page in list(doc)[:MAX_PAGES]:
        raw = page.get_text("text") or ""
        if len(raw.strip()) < 20:
            ocr_text = _ocr_page(page)
            if ocr_text.strip():
                raw = ocr_text
                ocr_used = True
            else:
                empty_pages += 1
        raw_pages.append(raw)
    headings, run_ins = styled_lines(list(doc)[:MAX_PAGES])
    page_texts = [reflow_pdf_lines(p, headings, run_ins) for p in strip_running_lines(raw_pages)]
    if empty_pages and not ocr_used:
        warnings.append(
            f"{empty_pages} page(s) had no extractable text (likely scanned). Install Tesseract + pytesseract to enable OCR."
        )
    text, pages = _join_pages(page_texts)
    hints: List[int] = []
    if headings:
        pos = 0
        for para in text.split("\n\n"):
            if para.strip() in headings:
                hints.append(pos + len(para) - len(para.lstrip()))
            pos += len(para) + 2
    return IngestedDocument(text=text, pages=pages, source_type="pdf", page_count=page_count,
                            ocr_used=ocr_used, heading_hints=hints, warnings=warnings)


def parse_docx(data: bytes) -> IngestedDocument:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentParseError("DOCX support requires python-docx.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"Could not open DOCX: {exc}") from exc

    blocks: List[Tuple[str, bool]] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            txt = para.text.strip()
            if not txt:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            is_heading = style.lower().startswith(("heading", "title")) or _is_heading_line(txt)
            blocks.append((txt, is_heading))
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    c = cell.text.strip()
                    if c and (not cells or cells[-1] != c):
                        cells.append(c)
                if len(cells) > 1 and _MARKER_CELL.match(cells[0]):
                    cells[:2] = [f"{cells[0]} {cells[1]}"]
                if cells:
                    blocks.append((" | ".join(cells), False))

    parts: List[str] = []
    hints: List[int] = []
    cursor = 0
    for txt, is_heading in blocks:
        txt = normalize_text(txt).replace("\n\n", "\n")
        if parts:
            parts.append("\n\n")
            cursor += 2
        if is_heading:
            hints.append(cursor)
        parts.append(txt)
        cursor += len(txt)
    text = "".join(parts)
    return IngestedDocument(text=text, pages=[PageSpan(page=1, start=0, end=len(text))], source_type="docx",
                            page_count=1, heading_hints=hints)


def parse_text(data: bytes) -> IngestedDocument:
    try:
        raw = data.decode("utf-8")
    except UnicodeDecodeError:
        raw = data.decode("latin-1", errors="replace")
    text = normalize_text(raw)
    return IngestedDocument(text=text, pages=[PageSpan(page=1, start=0, end=len(text))], source_type="txt", page_count=1)


def parse_document(data: bytes, filename: str, content_type: Optional[str] = None) -> IngestedDocument:
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        result = parse_pdf(data)
    elif name.endswith(".docx"):
        result = parse_docx(data)
    elif name.endswith(".doc"):
        raise DocumentParseError("Legacy .doc files are not supported. Save the file as .docx or PDF and retry.")
    else:
        if b"\x00" in data[:1024]:
            raise DocumentParseError("Unsupported binary file. Upload a PDF, DOCX or plain-text file.")
        result = parse_text(data)

    if len(result.text) > MAX_CHARS:
        result.text = result.text[:MAX_CHARS]
        result.pages = [p for p in result.pages if p.start < MAX_CHARS]
        if result.pages:
            result.pages[-1].end = min(result.pages[-1].end, MAX_CHARS)
        result.warnings.append(f"Document truncated to the first {MAX_CHARS:,} characters.")
    if len(result.text.strip()) < 40:
        raise DocumentParseError("Not enough readable text was found in this document to analyze.")
    return result


def page_for_offset(pages: List[PageSpan], offset: int) -> Optional[int]:
    for span in pages:
        if span.start <= offset <= span.end:
            return span.page
    return pages[-1].page if pages else None
